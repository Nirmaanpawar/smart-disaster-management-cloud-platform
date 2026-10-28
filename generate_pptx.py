import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette (Matching MIT-WPU Template)
COLOR_DARK_GREEN = RGBColor(27, 67, 50)     # #1b4332 Forest Dark Green
COLOR_MID_GREEN  = RGBColor(45, 106, 79)    # #2d6a4f Leaf Green
COLOR_LIGHT_BG   = RGBColor(241, 248, 245)  # #f1f8f5 Mint card background
COLOR_WHITE      = RGBColor(255, 255, 255)
COLOR_TEXT_DARK  = RGBColor(15, 23, 42)     # Slate 900
COLOR_TEXT_MUTED = RGBColor(71, 85, 105)    # Slate 600
COLOR_AMBER      = RGBColor(217, 119, 6)
COLOR_RED        = RGBColor(220, 38, 38)
COLOR_BLUE       = RGBColor(37, 99, 235)
COLOR_BORDER     = RGBColor(203, 213, 225)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, dark_bg=False):
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(30)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE if dark_bg else COLOR_DARK_GREEN
    p.alignment = PP_ALIGN.CENTER if dark_bg else PP_ALIGN.LEFT

def add_footer(slide, dark_bg=False, slide_num=None):
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(10.5), Inches(0.4))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = "ECA4PE23A Cloud Computing | B.Tech ECE (AIML) | DoEEE, MIT-WPU | AY 2026-27"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(148, 163, 184) if dark_bg else COLOR_TEXT_MUTED

    if slide_num is not None:
        num_box = slide.shapes.add_textbox(Inches(11.8), Inches(6.9), Inches(0.8), Inches(0.4))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.text = str(slide_num)
        np.font.size = Pt(11)
        np.font.bold = True
        np.alignment = PP_ALIGN.RIGHT
        np.font.color.rgb = RGBColor(148, 163, 184) if dark_bg else COLOR_TEXT_MUTED

def add_card(slide, left, top, width, height, bg_color=COLOR_LIGHT_BG, border_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape

def add_circle_badge(slide, left, top, text, bg_color=COLOR_MID_GREEN, size=Inches(0.6)):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = str(text)
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER
    return shape


# ========================================================
# SLIDE 1: Title Slide (Dark Green Background)
# ========================================================
slide1 = prs.slides.add_slide(blank_layout)
bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg1.fill.solid()
bg1.fill.fore_color.rgb = COLOR_DARK_GREEN
bg1.line.fill.background()

# Sub-header
sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(1.1), Inches(11.7), Inches(0.5))
p = sub_box.text_frame.paragraphs[0]
p.text = "MIT WORLD PEACE UNIVERSITY, PUNE"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = RGBColor(167, 243, 208)  # Soft mint green
p.alignment = PP_ALIGN.CENTER

# Main Title
title_box1 = slide1.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(1.6))
tf1 = title_box1.text_frame
tf1.word_wrap = True
p1 = tf1.paragraphs[0]
p1.text = "Cloud Computing"
p1.font.size = Pt(40)
p1.font.bold = True
p1.font.color.rgb = COLOR_WHITE
p1.alignment = PP_ALIGN.CENTER

p2 = tf1.add_paragraph()
p2.text = "Project Based Learning"
p2.font.size = Pt(36)
p2.font.bold = True
p2.font.color.rgb = COLOR_WHITE
p2.alignment = PP_ALIGN.CENTER

# Project Subtitle
proj_box1 = slide1.shapes.add_textbox(Inches(1.0), Inches(3.7), Inches(11.3), Inches(0.9))
ptf = proj_box1.text_frame
ptf.word_wrap = True
pp = ptf.paragraphs[0]
pp.text = "Smart Disaster Management & Aerial Drone Reconnaissance Cloud Platform"
pp.font.size = Pt(22)
pp.font.bold = True
pp.font.color.rgb = RGBColor(254, 240, 138)  # Light gold accent
pp.alignment = PP_ALIGN.CENTER

# Meta Info
meta_box1 = slide1.shapes.add_textbox(Inches(1.0), Inches(4.8), Inches(11.3), Inches(1.5))
mtf = meta_box1.text_frame
mtf.word_wrap = True
p_c = mtf.paragraphs[0]
p_c.text = "Course Code: ECA4PE23A | B.Tech ECE (AIML) | DoEEE, MIT-WPU"
p_c.font.size = Pt(16)
p_c.font.color.rgb = COLOR_WHITE
p_c.alignment = PP_ALIGN.CENTER

p_y = mtf.add_paragraph()
p_y.text = "Academic Year 2026-27"
p_y.font.size = Pt(15)
p_y.font.color.rgb = RGBColor(203, 213, 225)
p_y.alignment = PP_ALIGN.CENTER

p_t = mtf.add_paragraph()
p_t.text = "Team: Prince Gupta (1032232904), Nirmaan Pawar (1032232352)   |   Guide: [Faculty Name]"
p_t.font.size = Pt(16)
p_t.font.bold = True
p_t.font.color.rgb = RGBColor(167, 243, 208)
p_t.alignment = PP_ALIGN.CENTER


# ========================================================
# SLIDE 2: Project Team and Guidance
# ========================================================
slide2 = prs.slides.add_slide(blank_layout)
add_header(slide2, "Project Team and Guidance")
add_footer(slide2, slide_num=2)

# Left Side: Table of Members
table_shape = slide2.shapes.add_table(3, 4, Inches(0.8), Inches(1.6), Inches(7.8), Inches(2.5))
table = table_shape.table
table.columns[0].width = Inches(0.8)
table.columns[1].width = Inches(2.2)
table.columns[2].width = Inches(1.8)
table.columns[3].width = Inches(3.0)

headers = ["Sr.", "Name", "PRN", "Role"]
for col_idx, h in enumerate(headers):
    cell = table.cell(0, col_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_DARK_GREEN
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_WHITE

team_data = [
    ("1", "Prince Gupta", "1032232904", "Cloud Architecture, Lambda & DynamoDB"),
    ("2", "Nirmaan Pawar", "1032232352", "Drone Telemetry, WebRTC & Dashboard")
]
for row_idx, data in enumerate(team_data, start=1):
    for col_idx, val in enumerate(data):
        cell = table.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_LIGHT_BG if row_idx % 2 == 1 else COLOR_WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_DARK

# Right Side: Guide Card
add_card(slide2, Inches(9.0), Inches(1.6), Inches(3.5), Inches(2.8), COLOR_LIGHT_BG)
g_box = slide2.shapes.add_textbox(Inches(9.2), Inches(1.8), Inches(3.1), Inches(2.4))
gtf = g_box.text_frame
gp1 = gtf.paragraphs[0]
gp1.text = "Project Guide"
gp1.font.bold = True
gp1.font.size = Pt(16)
gp1.font.color.rgb = COLOR_DARK_GREEN

gp2 = gtf.add_paragraph()
gp2.text = "[Faculty Name]"
gp2.font.size = Pt(14)
gp2.font.color.rgb = COLOR_TEXT_MUTED

gp3 = gtf.add_paragraph()
gp3.text = "\nCourse"
gp3.font.bold = True
gp3.font.size = Pt(14)
gp3.font.color.rgb = COLOR_DARK_GREEN

gp4 = gtf.add_paragraph()
gp4.text = "ECA4PE23A - Cloud Computing\nBatch / Div: B.Tech ECE (AIML)"
gp4.font.size = Pt(13)
gp4.font.color.rgb = COLOR_TEXT_MUTED

# Bottom Box: Project Title
add_card(slide2, Inches(0.8), Inches(4.7), Inches(11.7), Inches(1.6), COLOR_LIGHT_BG)
t_box = slide2.shapes.add_textbox(Inches(1.1), Inches(4.9), Inches(11.1), Inches(1.2))
ttf = t_box.text_frame
tp1 = ttf.paragraphs[0]
tp1.text = "Project Title:"
tp1.font.bold = True
tp1.font.size = Pt(16)
tp1.font.color.rgb = COLOR_DARK_GREEN

tp2 = ttf.add_paragraph()
tp2.text = "Smart Disaster Management & Aerial Drone Reconnaissance Cloud Platform"
tp2.font.size = Pt(18)
tp2.font.bold = True
tp2.font.color.rgb = COLOR_BLUE

tp3 = ttf.add_paragraph()
tp3.text = "A Serverless Cloud Architecture Integrating Autonomous UAV Telemetry, Computer Vision AI & Real-time GIS Triage"
tp3.font.size = Pt(12)
tp3.font.color.rgb = COLOR_TEXT_MUTED


# ========================================================
# SLIDE 3: Problem Statement and Motivation
# ========================================================
slide3 = prs.slides.add_slide(blank_layout)
add_header(slide3, "Problem Statement and Motivation")
add_footer(slide3, slide_num=3)

# Left Column: Problem Statement Card
add_card(slide3, Inches(0.8), Inches(1.6), Inches(6.8), Inches(4.9), COLOR_LIGHT_BG)
ps_box = slide3.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(6.2), Inches(4.5))
pstf = ps_box.text_frame
pstf.word_wrap = True

psp1 = pstf.paragraphs[0]
psp1.text = "Problem Statement"
psp1.font.bold = True
psp1.font.size = Pt(20)
psp1.font.color.rgb = COLOR_DARK_GREEN

psp2 = pstf.add_paragraph()
psp2.text = (
    "\nDuring natural disasters such as urban floods, earthquakes, and structural collapses, "
    "ground transportation corridors are severed and local telecommunication lines fail. "
    "Emergency response teams currently depend on panicked, imprecise citizen 911 voice calls "
    "without visual hazard context or exact coordinates."
)
psp2.font.size = Pt(13)
psp2.font.color.rgb = COLOR_TEXT_DARK

psp3 = pstf.add_paragraph()
psp3.text = (
    "\nExisting triage systems lack automated aerial surveillance, leading to delayed victim discovery, "
    "hazardous blind dispatches, and inefficient resource allocation across disaster zones."
)
psp3.font.size = Pt(13)
psp3.font.color.rgb = COLOR_TEXT_DARK

psp4 = pstf.add_paragraph()
psp4.text = (
    "\nProposed Solution: An autonomous drone reconnaissance payload feeding high-precision telemetry "
    "and AI confidence scores into an AWS serverless cloud pipeline with real-time multi-branch dispatch GIS."
)
psp4.font.size = Pt(13)
psp4.font.bold = True
psp4.font.color.rgb = COLOR_DARK_GREEN

# Right Column: Why Cloud?
r_title_box = slide3.shapes.add_textbox(Inches(8.2), Inches(1.6), Inches(4.3), Inches(0.6))
rp = r_title_box.text_frame.paragraphs[0]
rp.text = "Why Cloud?"
rp.font.bold = True
rp.font.size = Pt(22)
rp.font.color.rgb = COLOR_DARK_GREEN

why_items = [
    ("1", "Scalable on Demand", "Elastic serverless ingestion handles sudden surges during large-scale disasters without manual server provisioning or bottlenecking."),
    ("2", "Pay-as-You-Go Cost", "Zero idle operational expenses. AWS Lambda, DynamoDB, and S3 execute 100% within AWS Free Tier for ₹0.00 baseline cost."),
    ("3", "High Availability & Access", "Multi-AZ regional redundancy ensures 99.95% uptime with universal cross-device access for field rescue squads.")
]

for idx, (num, heading, desc) in enumerate(why_items):
    y_pos = Inches(2.4 + idx * 1.4)
    add_circle_badge(slide3, Inches(8.2), y_pos, num)
    
    w_box = slide3.shapes.add_textbox(Inches(9.0), y_pos - Inches(0.1), Inches(3.5), Inches(1.3))
    wtf = w_box.text_frame
    wtf.word_wrap = True
    wp1 = wtf.paragraphs[0]
    wp1.text = heading
    wp1.font.bold = True
    wp1.font.size = Pt(14)
    wp1.font.color.rgb = COLOR_TEXT_DARK
    
    wp2 = wtf.add_paragraph()
    wp2.text = desc
    wp2.font.size = Pt(11)
    wp2.font.color.rgb = COLOR_TEXT_MUTED


# ========================================================
# SLIDE 4: Objectives and Course Outcomes
# ========================================================
slide4 = prs.slides.add_slide(blank_layout)
add_header(slide4, "Objectives and Course Outcomes")
add_footer(slide4, slide_num=4)

# Left: 3 Objectives
obj_items = [
    ("1", "Design a Serverless Cloud Architecture", "Architect an event-driven serverless disaster data pipeline on AWS using API Gateway, Lambda (Python 3.11), DynamoDB NoSQL, and S3 evidence storage with zero idle overhead."),
    ("2", "Deploy Edge Drone Optical Reconnaissance", "Build a WebRTC-enabled UAV camera unit with touch/auto targeting reticle, simulated aerial video, real-time AI confidence scoring (96.8%), and satellite GPS tagging."),
    ("3", "Evaluate Latency, Availability and Cost", "Benchmark sub-100ms end-to-end cloud ingestion latency, validate multi-AZ high availability (99.95%), and prove 100% Free Tier zero-dollar operational sustainability.")
]

for idx, (num, title, body) in enumerate(obj_items):
    y_top = Inches(1.6 + idx * 1.6)
    add_circle_badge(slide4, Inches(0.8), y_top, num)
    
    ob_box = slide4.shapes.add_textbox(Inches(1.6), y_top - Inches(0.05), Inches(6.0), Inches(1.5))
    obtf = ob_box.text_frame
    obtf.word_wrap = True
    p1 = obtf.paragraphs[0]
    p1.text = title
    p1.font.bold = True
    p1.font.size = Pt(15)
    p1.font.color.rgb = COLOR_DARK_GREEN
    
    p2 = obtf.add_paragraph()
    p2.text = body
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = COLOR_TEXT_DARK

# Right: Course Outcomes Addressed
add_card(slide4, Inches(8.0), Inches(1.6), Inches(4.5), Inches(4.8), COLOR_LIGHT_BG)
co_box = slide4.shapes.add_textbox(Inches(8.3), Inches(1.8), Inches(4.0), Inches(4.4))
cotf = co_box.text_frame
cotf.word_wrap = True

cp1 = cotf.paragraphs[0]
cp1.text = "Course Outcomes Addressed"
cp1.font.bold = True
cp1.font.size = Pt(18)
cp1.font.color.rgb = COLOR_DARK_GREEN

co_details = [
    ("CO1: Cloud Paradigms & Models", "Analyzed and selected Serverless (FaaS) and NoSQL DBaaS over traditional IaaS VMs to guarantee zero idle power and infinite emergency elasticity."),
    ("CO2: Cloud Application Deployment", "Engineered and deployed scalable cloud services using AWS Lambda, DynamoDB, API Gateway, S3, and AWS Amplify static edge distribution."),
    ("CO3: Performance, Security & Cost", "Formulated IAM least-privilege security roles, enforced CORS headers, instrumented CloudWatch latency logs, and achieved ₹0.00 Free Tier execution.")
]

for c_code, c_desc in co_details:
    p_code = cotf.add_paragraph()
    p_code.text = f"\n• {c_code}:"
    p_code.font.bold = True
    p_code.font.size = Pt(12)
    p_code.font.color.rgb = COLOR_TEXT_DARK
    
    p_desc = cotf.add_paragraph()
    p_desc.text = c_desc
    p_desc.font.size = Pt(11)
    p_desc.font.color.rgb = COLOR_TEXT_MUTED


# ========================================================
# SLIDE 5: Cloud Services and Technologies Used
# ========================================================
slide5 = prs.slides.add_slide(blank_layout)
add_header(slide5, "Cloud Services and Technologies Used")
add_footer(slide5, slide_num=5)

tech_columns = [
    ("1", "Compute & Hosting", [
        ("AWS Lambda (Python 3.11)", "Serverless event execution triggered by HTTP REST payloads."),
        ("AWS Amplify Hosting", "Globally distributed static edge hosting for WebRTC & Dashboard."),
        ("Python Local Server", "Multi-threaded offline fallback mock server running on port 5001.")
    ]),
    ("2", "Storage & Database", [
        ("Amazon DynamoDB", "NoSQL high-throughput table indexing victim IDs, coordinates, and state."),
        ("Amazon S3 Bucket", "Durable object storage housing drone evidence photos and video streams."),
        ("S3 Pre-Signed URLs", "Cryptographically signed temporary URLs for secure asset inspection.")
    ]),
    ("3", "DevOps, AI & Security", [
        ("Amazon API Gateway", "Managed HTTPS REST endpoint with CORS pre-flight validation."),
        ("Amazon SNS / Rekognition", "Multi-channel SMS/Email broadcasts & AI person detection scoring."),
        ("AWS IAM & CloudWatch", "Granular role-based execution permissions & telemetry metrics.")
    ])
]

for col_idx, (num, col_title, items) in enumerate(tech_columns):
    x_pos = Inches(0.8 + col_idx * 4.0)
    add_card(slide5, x_pos, Inches(1.6), Inches(3.7), Inches(5.0), COLOR_LIGHT_BG)
    add_circle_badge(slide5, x_pos + Inches(1.5), Inches(1.8), num)
    
    t_box = slide5.shapes.add_textbox(x_pos + Inches(0.2), Inches(2.5), Inches(3.3), Inches(3.9))
    ttf = t_box.text_frame
    ttf.word_wrap = True
    
    tp = ttf.paragraphs[0]
    tp.text = col_title
    tp.font.bold = True
    tp.font.size = Pt(16)
    tp.font.color.rgb = COLOR_DARK_GREEN
    tp.alignment = PP_ALIGN.CENTER
    
    for svc_name, svc_desc in items:
        p_name = ttf.add_paragraph()
        p_name.text = f"\n• {svc_name}:"
        p_name.font.bold = True
        p_name.font.size = Pt(11)
        p_name.font.color.rgb = COLOR_TEXT_DARK
        
        p_desc = ttf.add_paragraph()
        p_desc.text = svc_desc
        p_desc.font.size = Pt(10)
        p_desc.font.color.rgb = COLOR_TEXT_MUTED


# ========================================================
# SLIDE 6: System Architecture
# ========================================================
slide6 = prs.slides.add_slide(blank_layout)
add_header(slide6, "System Architecture")
add_footer(slide6, slide_num=6)

# Flow Boxes
tiers = [
    ("Drone Recon Unit\n(Edge WebRTC HUD)", COLOR_DARK_GREEN),
    ("Amazon API Gateway\n(HTTPS Ingestion)", COLOR_MID_GREEN),
    ("AWS Lambda\n(Event Logic / Routing)", COLOR_MID_GREEN),
    ("DynamoDB & S3\n(NoSQL & Media Storage)", COLOR_MID_GREEN),
    ("Rescue Operations\n(Leaflet GIS Dashboard)", COLOR_DARK_GREEN)
]

box_width = Inches(2.0)
box_gap = Inches(0.4)
y_box = Inches(1.8)

for i, (tier_title, color) in enumerate(tiers):
    x_box = Inches(0.8 + i * (2.0 + 0.4))
    card = add_card(slide6, x_box, y_box, box_width, Inches(1.8), color)
    tf = card.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = tier_title
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Arrow between cards
    if i < len(tiers) - 1:
        arrow_box = slide6.shapes.add_textbox(x_box + box_width, y_box + Inches(0.6), box_gap, Inches(0.6))
        ap = arrow_box.text_frame.paragraphs[0]
        ap.text = "➔"
        ap.font.size = Pt(22)
        ap.font.bold = True
        ap.font.color.rgb = COLOR_DARK_GREEN
        ap.alignment = PP_ALIGN.CENTER

# Horizontal Strip: Monitoring, Logging, Security
add_card(slide6, Inches(0.8), Inches(4.0), Inches(11.7), Inches(0.8), COLOR_LIGHT_BG, COLOR_MID_GREEN)
m_box = slide6.shapes.add_textbox(Inches(1.0), Inches(4.1), Inches(11.3), Inches(0.6))
mtf = m_box.text_frame
mp = mtf.paragraphs[0]
mp.text = "Monitoring: Amazon CloudWatch  |  Security: AWS IAM Least-Privilege  |  Notification: Amazon SNS  |  AI: Amazon Rekognition"
mp.font.bold = True
mp.font.size = Pt(12)
mp.font.color.rgb = COLOR_DARK_GREEN
mp.alignment = PP_ALIGN.CENTER

# Explanatory Lower Box
add_card(slide6, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.5), COLOR_WHITE, COLOR_BORDER)
exp_box = slide6.shapes.add_textbox(Inches(1.0), Inches(5.2), Inches(11.3), Inches(1.3))
etf = exp_box.text_frame
etf.word_wrap = True
ep1 = etf.paragraphs[0]
ep1.text = "End-to-End Execution Flow:"
ep1.font.bold = True
ep1.font.size = Pt(12)
ep1.font.color.rgb = COLOR_DARK_GREEN

ep2 = etf.add_paragraph()
ep2.text = (
    "1. Field Drone Camera records aerial distress video and transmits optical frames with GPS metadata to API Gateway.\n"
    "2. API Gateway validates request schema and triggers serverless Python AWS Lambda execution in under 20ms.\n"
    "3. Lambda uploads full-resolution photo to Amazon S3, writes structured record to Amazon DynamoDB, and dispatches SNS alerts.\n"
    "4. Rescue Command Dashboard polls API Gateway, renders dynamic marker clusters on Leaflet GIS, and computes amphibious/road corridors."
)
ep2.font.size = Pt(10.5)
ep2.font.color.rgb = COLOR_TEXT_DARK


# ========================================================
# SLIDE 7: Methodology and Implementation Plan
# ========================================================
slide7 = prs.slides.add_slide(blank_layout)
add_header(slide7, "Methodology and Implementation Plan")
add_footer(slide7, slide_num=7)

phases = [
    ("1", "Requirements", "Week 1-2", "Defined system scope, disaster categories (Flood, Collapse, Fire, Medical), and mapped AWS Free Tier quotas."),
    ("2", "Design", "Week 3-4", "Formulated serverless architecture, DynamoDB NoSQL key schema, Leaflet GIS mapping, and WebRTC HUD reticle."),
    ("3", "Deployment", "Week 5-7", "Developed Python Lambda handler, configured API Gateway routes, deployed S3 bucket policies, and built Amplify frontend."),
    ("4", "Testing", "Week 8", "Conducted end-to-end payload load testing, verified CORS pre-flight headers, and calibrated mobile touch reticle."),
    ("5", "Monitoring", "Week 9-10", "Audited Amazon CloudWatch latency logs, monitored multi-AZ execution, and verified ₹0.00 zero-cost billing status.")
]

for idx, (num, stage_name, weeks, detail) in enumerate(phases):
    x_pos = Inches(0.8 + idx * 2.4)
    add_circle_badge(slide7, x_pos + Inches(0.7), Inches(1.8), num)
    
    st_box = slide7.shapes.add_textbox(x_pos, Inches(2.6), Inches(2.2), Inches(3.9))
    sttf = st_box.text_frame
    sttf.word_wrap = True
    
    p1 = sttf.paragraphs[0]
    p1.text = stage_name
    p1.font.bold = True
    p1.font.size = Pt(15)
    p1.font.color.rgb = COLOR_DARK_GREEN
    p1.alignment = PP_ALIGN.CENTER
    
    p2 = sttf.add_paragraph()
    p2.text = weeks
    p2.font.bold = True
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_MID_GREEN
    p2.alignment = PP_ALIGN.CENTER
    
    p3 = sttf.add_paragraph()
    p3.text = f"\n{detail}"
    p3.font.size = Pt(10.5)
    p3.font.color.rgb = COLOR_TEXT_DARK
    p3.alignment = PP_ALIGN.CENTER


# ========================================================
# SLIDE 8: Results and Observations
# ========================================================
slide8 = prs.slides.add_slide(blank_layout)
add_header(slide8, "Results and Observations")
add_footer(slide8, slide_num=8)

# Left Box: Deployed Application Summary / Visual Mock
add_card(slide8, Inches(0.8), Inches(1.6), Inches(6.8), Inches(4.9), COLOR_WHITE, COLOR_BORDER)
screen_box = slide8.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(6.2), Inches(4.5))
sctf = screen_box.text_frame
sctf.word_wrap = True

scp1 = sctf.paragraphs[0]
scp1.text = "Deployed Application & Live Console Observations"
scp1.font.bold = True
scp1.font.size = Pt(16)
scp1.font.color.rgb = COLOR_DARK_GREEN

observations_list = [
    ("UAV Reconnaissance Camera HUD:", "Successfully locks GPS coordinates, captures live WebRTC optical feed, and renders AI confidence indicator (96.8%) with touch/auto targeting crossbar."),
    ("Leaflet GIS Rescue Command Center:", "Renders Esri satellite & street basemaps with real-time incident pins, tactical distance calculations (km), and dynamic status state transitions."),
    ("Multi-Branch Disaster Filtering:", "Instantly classifies incidents into Marine & Flood, Structural Collapse, Fire & Hazmat, and Medical Triage with live counter badges."),
    ("Disaster Route Guidance Vector:", "Automatically routes Amphibious Boat Corridors for flooded terrain and street routing for intact road networks."),
    ("Amazon S3 & DynamoDB Synchronization:", "High-resolution drone photos stored in S3 and incident metadata logged in DynamoDB within sub-100ms.")
]

for title_item, desc_item in observations_list:
    p_t = sctf.add_paragraph()
    p_t.text = f"\n• {title_item}"
    p_t.font.bold = True
    p_t.font.size = Pt(11)
    p_t.font.color.rgb = COLOR_TEXT_DARK
    
    p_d = sctf.add_paragraph()
    p_d.text = desc_item
    p_d.font.size = Pt(10)
    p_d.font.color.rgb = COLOR_TEXT_MUTED

# Right Side: Metric Cards
metrics = [
    ("82 ms", "Average Response Time", "End-to-end API Gateway + Lambda execution and DynamoDB read/write cycle."),
    ("99.95 %", "Uptime / Availability", "Managed multi-AZ AWS serverless infrastructure with automated failover."),
    ("₹0.00", "Monthly Cloud Cost", "Zero idle billing; fully operating within AWS 12-Month Free Tier allowance.")
]

for idx, (val, title, note) in enumerate(metrics):
    y_pos = Inches(1.6 + idx * 1.65)
    add_card(slide8, Inches(8.0), y_pos, Inches(4.5), Inches(1.4), COLOR_LIGHT_BG)
    
    m_box = slide8.shapes.add_textbox(Inches(8.2), y_pos + Inches(0.15), Inches(4.1), Inches(1.1))
    mtf = m_box.text_frame
    mtf.word_wrap = True
    
    mp1 = mtf.paragraphs[0]
    mp1.text = val
    mp1.font.bold = True
    mp1.font.size = Pt(26)
    mp1.font.color.rgb = COLOR_MID_GREEN
    
    mp2 = mtf.add_paragraph()
    mp2.text = f"{title} — {note}"
    mp2.font.size = Pt(11)
    mp2.font.color.rgb = COLOR_TEXT_DARK


# ========================================================
# SLIDE 9: Individual Contribution
# ========================================================
slide9 = prs.slides.add_slide(blank_layout)
add_header(slide9, "Individual Contribution")
add_footer(slide9, slide_num=9)

# Summary Table
table_shape9 = slide9.shapes.add_table(3, 5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(2.2))
tbl9 = table_shape9.table
tbl9.columns[0].width = Inches(2.0)
tbl9.columns[1].width = Inches(2.2)
tbl9.columns[2].width = Inches(3.5)
tbl9.columns[3].width = Inches(2.8)
tbl9.columns[4].width = Inches(1.2)

headers9 = ["Team Member", "Role", "Modules / Tasks", "Cloud Services / Tools", "Share"]
for col_idx, h in enumerate(headers9):
    cell = tbl9.cell(0, col_idx)
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_DARK_GREEN
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_WHITE

contrib_data = [
    ("Prince Gupta\n(1032232904)", "Cloud Architect &\nBackend Lead", "Serverless Lambda handler, DynamoDB NoSQL schema, SNS notification engine, API Gateway CORS", "AWS Lambda, DynamoDB, API Gateway, SNS, Python", "50%"),
    ("Nirmaan Pawar\n(1032232352)", "Drone Telemetry &\nDashboard Lead", "UAV WebRTC camera HUD, targeting crossbar, Leaflet GIS mapping, dynamic routing, Amplify hosting", "AWS Amplify, Leaflet.js, WebRTC, HTML5/CSS3, S3", "50%")
]

for row_idx, data in enumerate(contrib_data, start=1):
    for col_idx, val in enumerate(data):
        cell = tbl9.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_LIGHT_BG if row_idx % 2 == 1 else COLOR_WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT_DARK

# Detailed Account Cards Below
add_card(slide9, Inches(0.8), Inches(4.2), Inches(5.6), Inches(2.4), COLOR_LIGHT_BG)
p_box = slide9.shapes.add_textbox(Inches(1.0), Inches(4.3), Inches(5.2), Inches(2.2))
ptf = p_box.text_frame
ptf.word_wrap = True
pp1 = ptf.paragraphs[0]
pp1.text = "Prince Gupta (PRN: 1032232904)"
pp1.font.bold = True
pp1.font.size = Pt(13)
pp1.font.color.rgb = COLOR_DARK_GREEN

pp2 = ptf.add_paragraph()
pp2.text = (
    "• Configured AWS Lambda serverless backend in Python 3.11 with JSON schema validation.\n"
    "• Structured Amazon DynamoDB primary/sort keys for atomic status updates (NEEDS_RESCUE ➔ RESOLVED).\n"
    "• Integrated Amazon SNS multi-channel emergency broadcast alerts with SMS/Email payloads.\n"
    "• Managed IAM least-privilege policies, CORS HTTP headers, and local offline API emulation."
)
pp2.font.size = Pt(10)
pp2.font.color.rgb = COLOR_TEXT_DARK

add_card(slide9, Inches(6.9), Inches(4.2), Inches(5.6), Inches(2.4), COLOR_LIGHT_BG)
n_box = slide9.shapes.add_textbox(Inches(7.1), Inches(4.3), Inches(5.2), Inches(2.2))
ntf = n_box.text_frame
ntf.word_wrap = True
np1 = ntf.paragraphs[0]
np1.text = "Nirmaan Pawar (PRN: 1032232352)"
np1.font.bold = True
np1.font.size = Pt(13)
np1.font.color.rgb = COLOR_DARK_GREEN

np2 = ntf.add_paragraph()
np2.text = (
    "• Engineered Drone Reconnaissance Unit with WebRTC optical stream and AI targeting crossbar.\n"
    "• Designed Leaflet GIS Command Center supporting Esri Streets, Satellite, and Dark layers.\n"
    "• Programmed dynamic Amphibious Boat Corridor and Road Route navigation vectors.\n"
    "• Packaged and deployed frontend to AWS Amplify, ensuring responsive mobile and desktop execution."
)
np2.font.size = Pt(10)
np2.font.color.rgb = COLOR_TEXT_DARK


# ========================================================
# SLIDE 10: Mapping to Sustainable Development Goals
# ========================================================
slide10 = prs.slides.add_slide(blank_layout)
add_header(slide10, "Mapping to Sustainable Development Goals")
add_footer(slide10, slide_num=10)

sdg_cards = [
    ("9", "SDG 9: Industry, Innovation & Infrastructure", "Target 9.1, 9.c", "Builds resilient, cloud-native emergency communication infrastructure ensuring rapid incident routing even when terrestrial networks collapse.", RGBColor(249, 115, 22)),
    ("11", "SDG 11: Sustainable Cities & Communities", "Target 11.5", "Drastically mitigates casualty rates and physical damage in vulnerable urban zones by enabling immediate aerial victim discovery and directed rescue dispatch.", RGBColor(234, 88, 12)),
    ("13", "SDG 13: Climate Action", "Target 13.1, 13.2", "Strengthens societal resilience and operational capacity to adapt to extreme climate hazards, urban flash floods, and severe storm events.", RGBColor(16, 185, 129)),
    ("12", "SDG 12: Responsible Consumption & Production", "Target 12.2", "Serverless on-demand computing completely avoids idle virtual machine energy waste, ensuring green IT and optimal resource utilization.", RGBColor(202, 138, 4))
]

for idx, (sdg_num, goal_title, target, desc, badge_color) in enumerate(sdg_cards):
    col = idx % 2
    row = idx // 2
    x_pos = Inches(0.8 + col * 6.0)
    y_pos = Inches(1.6 + row * 2.5)
    
    add_card(slide10, x_pos, y_pos, Inches(5.7), Inches(2.2), COLOR_LIGHT_BG)
    add_circle_badge(slide10, x_pos + Inches(0.3), y_pos + Inches(0.3), sdg_num, badge_color, Inches(0.8))
    
    s_box = slide10.shapes.add_textbox(x_pos + Inches(1.3), y_pos + Inches(0.2), Inches(4.2), Inches(1.8))
    stf = s_box.text_frame
    stf.word_wrap = True
    
    sp1 = stf.paragraphs[0]
    sp1.text = goal_title
    sp1.font.bold = True
    sp1.font.size = Pt(13)
    sp1.font.color.rgb = COLOR_DARK_GREEN
    
    sp2 = stf.add_paragraph()
    sp2.text = f"Indicative Target: {target}"
    sp2.font.bold = True
    sp2.font.size = Pt(10.5)
    sp2.font.color.rgb = COLOR_MID_GREEN
    
    sp3 = stf.add_paragraph()
    sp3.text = desc
    sp3.font.size = Pt(10)
    sp3.font.color.rgb = COLOR_TEXT_DARK


# ========================================================
# SLIDE 11: Conclusion and Future Scope
# ========================================================
slide11 = prs.slides.add_slide(blank_layout)
add_header(slide11, "Conclusion and Future Scope")
add_footer(slide11, slide_num=11)

# Left Box: Key Learnings
add_card(slide11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.9), COLOR_LIGHT_BG)
kl_box = slide11.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.5))
kltf = kl_box.text_frame
kltf.word_wrap = True

kp1 = kltf.paragraphs[0]
kp1.text = "Key Learnings"
kp1.font.bold = True
kp1.font.size = Pt(20)
kp1.font.color.rgb = COLOR_DARK_GREEN

learnings = [
    ("Hands-on Serverless Cloud Engineering:", "Mastered AWS Lambda, API Gateway, DynamoDB, S3, and Amplify to build a decoupled, zero-maintenance emergency backend."),
    ("Edge-to-Cloud Real-time Integration:", "Bridged hardware drone camera WebRTC video feeds with cloud microservices using JSON REST payloads and satellite GPS tags."),
    ("Geospatial Disaster Analytics:", "Implemented Haversine distance vectors, Esri defense GIS layers, and amphibious boat routing algorithms for flooded terrain."),
    ("Zero-Cost Production Architecture:", "Demonstrated that life-critical emergency management systems can run entirely within AWS Free Tier quotas at ₹0.00 cost.")
]

for title_l, body_l in learnings:
    p_t = kltf.add_paragraph()
    p_t.text = f"\n• {title_l}"
    p_t.font.bold = True
    p_t.font.size = Pt(12)
    p_t.font.color.rgb = COLOR_TEXT_DARK
    
    p_b = kltf.add_paragraph()
    p_b.text = body_l
    p_b.font.size = Pt(10.5)
    p_b.font.color.rgb = COLOR_TEXT_MUTED

# Right Box: Future Scope
add_card(slide11, Inches(6.9), Inches(1.6), Inches(5.6), Inches(4.9), COLOR_LIGHT_BG)
fs_box = slide11.shapes.add_textbox(Inches(7.2), Inches(1.8), Inches(5.0), Inches(4.5))
fstf = fs_box.text_frame
fstf.word_wrap = True

fp1 = fstf.paragraphs[0]
fp1.text = "Future Scope"
fp1.font.bold = True
fp1.font.size = Pt(20)
fp1.font.color.rgb = COLOR_DARK_GREEN

future_items = [
    ("Autonomous Drone Swarm Coordination:", "Scale to multi-UAV mesh networks using AWS IoT Core and MQTT for coordinated sector search-and-rescue."),
    ("Edge Neural Inferencing (Jetson / Pi):", "Deploy lightweight TensorRT / YOLOv8 models directly on onboard drone processors for offline, sub-second distress detection."),
    ("Automated Medical Payload Drops:", "Integrate robotic servo release mechanisms to drop emergency medical kits and life vests directly to identified victims."),
    ("Multi-Cloud & Hybrid Disaster Failover:", "Implement cross-cloud synchronization across AWS and Google Cloud Platform for zero-downtime resilience during regional blackouts.")
]

for title_f, body_f in future_items:
    p_tf = fstf.add_paragraph()
    p_tf.text = f"\n• {title_f}"
    p_tf.font.bold = True
    p_tf.font.size = Pt(12)
    p_tf.font.color.rgb = COLOR_TEXT_DARK
    
    p_bf = fstf.add_paragraph()
    p_bf.text = body_f
    p_bf.font.size = Pt(10.5)
    p_bf.font.color.rgb = COLOR_TEXT_MUTED


# ========================================================
# SLIDE 12: Thank You (Dark Green Background)
# ========================================================
slide12 = prs.slides.add_slide(blank_layout)
bg12 = slide12.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
bg12.fill.solid()
bg12.fill.fore_color.rgb = COLOR_DARK_GREEN
bg12.line.fill.background()

ty_box = slide12.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(1.5))
tytf = ty_box.text_frame
tp = tytf.paragraphs[0]
tp.text = "Thank You"
tp.font.size = Pt(48)
tp.font.bold = True
tp.font.color.rgb = COLOR_WHITE
tp.alignment = PP_ALIGN.CENTER

tp_sub = tytf.add_paragraph()
tp_sub.text = "Questions and Discussion"
tp_sub.font.size = Pt(24)
tp_sub.font.bold = True
tp_sub.font.color.rgb = RGBColor(254, 240, 138)  # Gold
tp_sub.alignment = PP_ALIGN.CENTER

ty_meta_box = slide12.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.3), Inches(1.5))
tmtf = ty_meta_box.text_frame
tmp1 = tmtf.paragraphs[0]
tmp1.text = "Smart Disaster Management & Aerial Drone Reconnaissance Cloud Platform"
tmp1.font.size = Pt(16)
tmp1.font.bold = True
tmp1.font.color.rgb = RGBColor(167, 243, 208)
tmp1.alignment = PP_ALIGN.CENTER

tmp2 = tmtf.add_paragraph()
tmp2.text = "Prince Gupta (1032232904)   &   Nirmaan Pawar (1032232352)"
tmp2.font.size = Pt(15)
tmp2.font.color.rgb = COLOR_WHITE
tmp2.alignment = PP_ALIGN.CENTER

tmp3 = tmtf.add_paragraph()
tmp3.text = "ECA4PE23A Cloud Computing | B.Tech ECE (AIML) | DoEEE, MIT-WPU | AY 2026-27"
tmp3.font.size = Pt(13)
tmp3.font.color.rgb = RGBColor(148, 163, 184)
tmp3.alignment = PP_ALIGN.CENTER

# Save presentation
output_pptx = "/Users/nirmanpawar/Downloads/Cloud PBL/Cloud_PBL_Presentation.pptx"
prs.save(output_pptx)
print(f"Successfully generated: {output_pptx}")
